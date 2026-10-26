"""Import ZIDAN DIVERS FIN 26 into Main Depot. Not a migration — one-shot admin RPC."""
from __future__ import annotations

import json
import os
import uuid
from collections import defaultdict
from io import BytesIO
from pathlib import Path

import urllib.error
import urllib.request
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "ZIDAN - DIVERS FIN 26.xlsx"
ENV = ROOT / ".env"

PLACEHOLDER_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
)


def load_env() -> dict[str, str]:
    out: dict[str, str] = {}
    for line in ENV.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        out[k.strip()] = v.strip()
    return out


def cell(ws, r, c):
    return ws.cell(r, c).value


def image_bytes(img) -> tuple[bytes, str]:
    data = img._data()
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return data, "png"
    if data[:2] == b"\xff\xd8":
        return data, "jpg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return data, "webp"
    return PLACEHOLDER_PNG, "png"


MAX_PHOTO_BYTES = 5 * 1024 * 1024
COMPRESS_ABOVE_BYTES = 350 * 1024
MAX_EDGE = 1600


def compress_photo(content: bytes, ext: str) -> tuple[bytes, str]:
    if len(content) <= COMPRESS_ABOVE_BYTES and len(content) <= MAX_PHOTO_BYTES:
        return content, ext
    img = Image.open(BytesIO(content))
    img = img.convert("RGB")
    w, h = img.size
    scale = min(1.0, MAX_EDGE / float(max(w, h)))
    if scale < 1.0:
        img = img.resize(
            (max(1, round(w * scale)), max(1, round(h * scale))),
            Image.Resampling.LANCZOS,
        )
    quality = 82
    out = b""
    while quality >= 40:
        buf = BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        out = buf.getvalue()
        if len(out) <= MAX_PHOTO_BYTES:
            print("COMPRESSED", len(content), "->", len(out), "q", quality)
            return out, "jpg"
        quality -= 10
    raise RuntimeError("photo still too large after compress: %s" % len(out))


def parse_workbook(ws):
    images_by_row: dict[int, list] = defaultdict(list)
    for img in getattr(ws, "_images", []) or []:
        try:
            row = img.anchor._from.row + 1
        except Exception:
            continue
        images_by_row[row].append(image_bytes(img))

    shop = None
    products = []
    pending = None
    for r in range(3, 125):
        a = cell(ws, r, 1)
        b = cell(ws, r, 2)
        d = cell(ws, r, 4)
        e = cell(ws, r, 5)
        f = cell(ws, r, 6)
        g = cell(ws, r, 7)
        h = cell(ws, r, 8)
        i = cell(ws, r, 9)
        if a:
            shop = str(a).replace("\n", " / ").strip()
        item = str(b).strip() if b else ""
        name = " ".join(str(d).split()) if d else ""

        if not item and not name:
            if pending and isinstance(e, (int, float)):
                pending["ctn"] = (pending["ctn"] or 0) + e
                extra_qty = g if isinstance(g, (int, float)) else e * (pending["pcs_ctn"] or 1)
                pending["tqty"] = (pending["tqty"] or 0) + extra_qty
                if images_by_row.get(r):
                    pending["photos"].extend(images_by_row[r])
            continue

        if not item and name.upper() == "ACCESSORY":
            continue

        rec = {
            "row": r,
            "shop": shop,
            "item": item,
            "name": name,
            "ctn": e if isinstance(e, (int, float)) else None,
            "pcs_ctn": f if isinstance(f, (int, float)) else None,
            "tqty": g if isinstance(g, (int, float)) else None,
            "cost": float(h) if isinstance(h, (int, float)) else None,
            "sell": float(i) if isinstance(i, (int, float)) else None,
            "photos": list(images_by_row.get(r, [])),
        }
        products.append(rec)
        pending = rec

    used: set[str] = set()
    for rec in products:
        serial = rec["item"]
        key = serial.lower()
        if key in used:
            suffix = "B22" if "B22" in rec["name"].upper() else "2"
            serial = "%s-%s" % (rec["item"], suffix)
            key = serial.lower()
        used.add(key)
        rec["serial"] = serial

        if rec["tqty"] is None and rec["ctn"] and rec["pcs_ctn"]:
            rec["tqty"] = rec["ctn"] * rec["pcs_ctn"]
        if rec["ctn"] is None and rec["tqty"] and rec["pcs_ctn"]:
            rec["ctn"] = rec["tqty"] / rec["pcs_ctn"]

        if rec["cost"] is not None and rec["sell"] is not None and rec["cost"] == rec["sell"]:
            rec["sell"] = 0.0
            rec["sell_pending"] = True
        else:
            rec["sell_pending"] = False

        rec["pcs_ctn"] = rec["pcs_ctn"] or 1
        rec["tqty"] = rec["tqty"] or 0

    last_photos = None
    for rec in products:
        if rec["photos"]:
            last_photos = rec["photos"]
            rec["photo_inherited"] = False
        elif last_photos:
            rec["photos"] = list(last_photos)
            rec["photo_inherited"] = True
        else:
            rec["photo_inherited"] = False

    return products


class Api:
    def __init__(self, url: str, anon: str, token: str):
        self.url = url.rstrip("/")
        self.anon = anon
        self.token = token

    def _headers(self, extra=None):
        h = {
            "apikey": self.anon,
            "authorization": "Bearer %s" % self.token,
            "content-type": "application/json",
        }
        if extra:
            h.update(extra)
        return h

    def request(self, method, path, data=None, extra=None):
        body = None if data is None else json.dumps(data).encode("utf-8")
        req = urllib.request.Request(
            self.url + path,
            data=body,
            method=method,
            headers=self._headers(extra),
        )
        try:
            with urllib.request.urlopen(req) as resp:
                payload = resp.read()
                if not payload:
                    return None
                return json.loads(payload.decode("utf-8"))
        except urllib.error.HTTPError as e:
            err = e.read().decode("utf-8", errors="replace")
            raise RuntimeError("%s %s -> %s %s" % (method, path, e.code, err)) from e

    def rpc(self, name, params):
        return self.request("POST", "/rest/v1/rpc/%s" % name, params)

    def select(self, table, query):
        return self.request(
            "GET",
            "/rest/v1/%s?%s" % (table, query),
            extra={"accept": "application/json"},
        )

    def upload_photo(self, content: bytes, ext: str) -> str:
        content, ext = compress_photo(content, ext)
        name = "%s.%s" % (uuid.uuid4(), ext)
        content_type = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}[ext]
        req = urllib.request.Request(
            self.url + "/storage/v1/object/product-photos/" + name,
            data=content,
            method="POST",
            headers={
                "apikey": self.anon,
                "authorization": "Bearer %s" % self.token,
                "content-type": content_type,
                "x-upsert": "false",
            },
        )
        try:
            with urllib.request.urlopen(req) as resp:
                resp.read()
        except urllib.error.HTTPError as e:
            err = e.read().decode("utf-8", errors="replace")
            raise RuntimeError("upload %s -> %s %s" % (name, e.code, err)) from e
        return "%s/storage/v1/object/public/product-photos/%s" % (self.url, name)


def login(url, anon, password):
    body = json.dumps({"email": "admin@vhm-pos.internal", "password": password}).encode("utf-8")
    req = urllib.request.Request(
        url.rstrip("/") + "/auth/v1/token?grant_type=password",
        data=body,
        method="POST",
        headers={"apikey": anon, "content-type": "application/json"},
    )
    with urllib.request.urlopen(req) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return data["access_token"]


def units_for(rec):
    cost = round(rec["cost"] or 0, 2)
    sell = round(rec["sell"] or 0, 2)
    pcs = int(rec["pcs_ctn"] or 1)
    piece = {
        "unit_level": 0,
        "unit_name": "Piece",
        "serial_number": rec["serial"],
        "relative_quantity": None,
        "cost_price_cfa": cost,
        "selling_price_cfa": sell,
    }
    if pcs <= 1:
        return [piece]
    return [
        piece,
        {
            "unit_level": 1,
            "unit_name": "Carton",
            "serial_number": rec["serial"] + "-CTN",
            "relative_quantity": pcs,
            "cost_price_cfa": round(cost * pcs, 2),
            "selling_price_cfa": round(sell * pcs, 2),
        },
    ]


def product_id_from(created):
    if isinstance(created, dict) and "id" in created:
        return created["id"]
    if isinstance(created, list) and created and isinstance(created[0], dict):
        return created[0]["id"]
    raise RuntimeError("Unexpected create_product payload: %r" % created)


def main():
    env = load_env()
    url = env["VITE_SUPABASE_URL"]
    anon = env["VITE_SUPABASE_ANON_KEY"]

    wb = load_workbook(XLSX, data_only=True)
    products = parse_workbook(wb.active)
    print("PARSED", len(products))
    print("SELL_PENDING", [p["serial"] for p in products if p["sell_pending"]])
    print("SERIAL_CHANGED", [(p["item"], p["serial"]) for p in products if p["item"] != p["serial"]])
    print("NO_PHOTO_SOURCE", [p["serial"] for p in products if not p["photos"]])
    print("SUM_PIECES", sum(p["tqty"] or 0 for p in products))

    if os.environ.get("DRY_RUN") == "1":
        return

    password = os.environ.get("ADMIN_PASSWORD")
    if not password:
        raise SystemExit("Set ADMIN_PASSWORD")

    token = login(url, anon, password)
    api = Api(url, anon, token)

    depots = api.select("depots", "select=id,name,code,is_active&order=created_at")
    main_depot = next(
        (
            d
            for d in depots
            if str(d.get("code", "")).upper() in ("MAIN", "MAIN_DEPOT")
            or "main" in str(d.get("name", "")).lower()
        ),
        depots[0] if depots else None,
    )
    if not main_depot:
        raise SystemExit("No main depot")
    print("USING_DEPOT", main_depot)

    existing_units = api.select("product_units", "select=serial_number") or []
    existing_serials = {str(u["serial_number"]).lower() for u in existing_units}

    cats = api.select("categories", "select=id,name,is_active") or []
    divers = next((c for c in cats if str(c.get("name", "")).strip().lower() == "divers"), None)
    if not divers:
        divers = api.rpc("create_category", {"p_name": "Divers", "p_parent_id": None})
    cat_id = divers["id"]
    print("CATEGORY", cat_id)

    report = []
    skipped = 0
    for rec in products:
        if rec["serial"].lower() in existing_serials:
            skipped += 1
            print("SKIP", rec["serial"])
            continue
        if rec["photos"]:
            raw, ext = rec["photos"][0]
        else:
            raw, ext = PLACEHOLDER_PNG, "png"
        photo_url = api.upload_photo(raw, ext)
        created = api.rpc(
            "create_product",
            {
                "p_depot_id": main_depot["id"],
                "p_category_id": cat_id,
                "p_name": rec["name"],
                "p_photo_url": photo_url,
                "p_units": units_for(rec),
                "p_track_low_stock": False,
                "p_low_stock_threshold": None,
            },
        )
        pid = product_id_from(created)
        restock = None
        if rec["tqty"] and rec["tqty"] > 0:
            restock = api.rpc(
                "restock_product",
                {
                    "p_product_id": pid,
                    "p_unit_level": 0,
                    "p_quantity_in_unit": rec["tqty"],
                    "p_note": "Opening stock — ZIDAN DIVERS FIN 26",
                },
            )
        existing_serials.add(rec["serial"].lower())
        report.append(
            {
                "serial": rec["serial"],
                "name": rec["name"],
                "product_id": pid,
                "tqty": rec["tqty"],
                "stock": None if restock is None else restock.get("new_stock_pieces"),
                "sell_pending": rec["sell_pending"],
            }
        )
        print("OK", rec["serial"], rec["tqty"])

    out = ROOT / "scripts" / "_zidan_import_report.json"
    out.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print("DONE new=%s skipped=%s wrote=%s" % (len(report), skipped, out))


if __name__ == "__main__":
    main()
