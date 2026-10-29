from __future__ import annotations
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "2INV.xlsx"
OUT = ROOT / "scripts" / "corrections" / "20261028_electro2_stock_cost.sql"
FX = 156

def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["ELECTRO 2"]
    rows = []
    for r in range(1, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        if not raw:
            continue
        serial = str(raw).strip()
        pcs = max(1, int(float(ws.cell(r, 5).value or 1)))
        tqty = int(float(ws.cell(r, 6).value or 0))
        c11 = float(ws.cell(r, 11).value or 0)
        l0 = round((c11 / pcs) * FX, 2)
        l1 = round(l0 * pcs, 2)
        rows.append((serial, pcs, tqty, l0, l1))
    values = ",\n".join(
        "    ('%s', %s, %s::numeric, %s::numeric, %s::numeric)" % (s, pcs, tqty, l0, l1)
        for s, pcs, tqty, l0, l1 in rows
    )
    sheet_value = sum(tqty * l0 for _, _, tqty, l0, _ in rows)
    sql = """-- NOT A MIGRATION. Outside supabase/migrations — `supabase db push` will not apply this.
-- ELECTRO 2 stock + cost correction (%(n)s products from 2INV.xlsx).
--
-- Bug (scripts/import_2inv_main.py parse_electro2): used C7 as pcs_ctn, C8 as tqty,
-- C13 as L0 cost. Truth: C4=CTN, C5=PCS/CTN, C6=TQTY pieces; C7/C8 are not pieces;
-- C13 = 126xC7 (not a cost). L0 cost = (C11 / C5) x FX; L1 cost = L0 x C5.
--
-- FX = %(fx)s CFA per foreign unit — CONFIRM before setting approved := true.
-- With FX=%(fx)s, Excel value at cost Sum(C6xL0) approx %(sv).0f CFA (not ~160 M).
-- If ~160 M was expected, FX or the C11 basis needs revisiting (C12x%(fx)s as L0 approx 198 M).
--
-- Stock target: GREATEST(0, C6 + Sum inventory_transactions.quantity_pieces
--   excluding RESTOCK) — Excel opening minus net sold/returned/transferred/adjusted.
-- Idempotent: skips products that already have audit action correct_electro2_stock_cost.
-- Refuses unless approved := true.

\\set ON_ERROR_STOP on

-- Before (preview; safe)
WITH truth(serial, pcs_ctn, tqty, l0_cost, l1_cost) AS (
  VALUES
%(values)s
),
targets AS (
  SELECT
    t.*,
    pu.product_id,
    p.depot_id,
    p.name AS product_name,
    p.current_stock_pieces AS stock_before,
    pu0.cost_price_cfa AS l0_before,
    pu1.cost_price_cfa AS l1_before,
    pu1.relative_quantity AS l1_rel_before,
    pu1.conversion_to_base AS l1_conv_before,
    GREATEST(
      0,
      t.tqty + COALESCE((
        SELECT SUM(it.quantity_pieces)
        FROM public.inventory_transactions it
        WHERE it.product_id = pu.product_id
          AND it.transaction_type <> 'RESTOCK'
      ), 0)
    ) AS stock_after
  FROM truth t
  JOIN public.product_units pu
    ON lower(pu.serial_number) = lower(t.serial) AND pu.unit_level = 0
  JOIN public.products p ON p.id = pu.product_id
  LEFT JOIN public.product_units pu0
    ON pu0.product_id = p.id AND pu0.unit_level = 0
  LEFT JOIN public.product_units pu1
    ON pu1.product_id = p.id AND pu1.unit_level = 1
)
SELECT
  serial, product_name, stock_before, stock_after,
  l0_before, l0_cost AS l0_after,
  l1_before, l1_cost AS l1_after,
  l1_rel_before, pcs_ctn AS l1_rel_after
FROM targets
ORDER BY serial;

-- Apply (gated)
DO $$
DECLARE
  approved boolean := false;  -- set true only after explicit approval
  r RECORD;
  v_stock_after numeric;
  v_before jsonb;
  v_after jsonb;
  v_done int := 0;
  v_skipped int := 0;
BEGIN
  IF NOT approved THEN
    RAISE EXCEPTION
      'ELECTRO 2 correction not approved. Set approved := true after reviewing the before SELECT.';
  END IF;

  FOR r IN
    WITH truth(serial, pcs_ctn, tqty, l0_cost, l1_cost) AS (
      VALUES
%(values)s
    )
    SELECT
      t.serial,
      t.pcs_ctn,
      t.tqty,
      t.l0_cost,
      t.l1_cost,
      pu.product_id,
      p.depot_id,
      p.current_stock_pieces AS stock_before,
      pu0.id AS l0_id,
      pu0.cost_price_cfa AS l0_before,
      pu0.conversion_to_base AS l0_conv_before,
      pu1.id AS l1_id,
      pu1.cost_price_cfa AS l1_before,
      pu1.relative_quantity AS l1_rel_before,
      pu1.conversion_to_base AS l1_conv_before
    FROM truth t
    JOIN public.product_units pu
      ON lower(pu.serial_number) = lower(t.serial) AND pu.unit_level = 0
    JOIN public.products p ON p.id = pu.product_id
    LEFT JOIN public.product_units pu0
      ON pu0.product_id = p.id AND pu0.unit_level = 0
    LEFT JOIN public.product_units pu1
      ON pu1.product_id = p.id AND pu1.unit_level = 1
  LOOP
    IF EXISTS (
      SELECT 1 FROM public.audit_log a
      WHERE a.entity_type = 'product'
        AND a.entity_id = r.product_id
        AND a.action = 'correct_electro2_stock_cost'
    ) THEN
      v_skipped := v_skipped + 1;
      CONTINUE;
    END IF;

    SELECT GREATEST(
      0,
      r.tqty + COALESCE(SUM(it.quantity_pieces), 0)
    )
    INTO v_stock_after
    FROM public.inventory_transactions it
    WHERE it.product_id = r.product_id
      AND it.transaction_type <> 'RESTOCK';

    v_before := jsonb_build_object(
      'serial', r.serial,
      'stock', r.stock_before,
      'l0_cost', r.l0_before,
      'l0_conversion', r.l0_conv_before,
      'l1_cost', r.l1_before,
      'l1_relative_quantity', r.l1_rel_before,
      'l1_conversion', r.l1_conv_before
    );

    UPDATE public.products
    SET current_stock_pieces = v_stock_after,
        updated_at = now()
    WHERE id = r.product_id;

    IF v_stock_after <> r.stock_before THEN
      INSERT INTO public.inventory_transactions (
        depot_id, product_id, quantity_pieces,
        transaction_type, reference_type, reference_id,
        created_by, metadata
      ) VALUES (
        r.depot_id, r.product_id, (v_stock_after - r.stock_before),
        'ADJUSTMENT', 'inventory_adjustment', r.product_id,
        (SELECT id FROM public.staff WHERE role = 'admin' AND is_active = true ORDER BY created_at LIMIT 1),
        jsonb_build_object(
          'source', 'correct_electro2_stock',
          'stock_before', r.stock_before,
          'stock_after', v_stock_after
        )
      );
    END IF;

    UPDATE public.product_units
    SET cost_price_cfa = r.l0_cost,
        conversion_to_base = 1,
        relative_quantity = NULL
    WHERE id = r.l0_id;

    IF r.l1_id IS NOT NULL THEN
      UPDATE public.product_units
      SET relative_quantity = r.pcs_ctn,
          conversion_to_base = r.pcs_ctn,
          cost_price_cfa = r.l1_cost
      WHERE id = r.l1_id;
    ELSIF r.pcs_ctn > 1 THEN
      INSERT INTO public.product_units (
        product_id, unit_level, unit_name, serial_number,
        relative_quantity, conversion_to_base,
        cost_price_cfa, selling_price_cfa
      ) VALUES (
        r.product_id, 1, 'Carton', r.serial || '-CTN',
        r.pcs_ctn, r.pcs_ctn,
        r.l1_cost,
        COALESCE(
          (SELECT selling_price_cfa * r.pcs_ctn FROM public.product_units WHERE id = r.l0_id),
          0
        )
      );
    END IF;

    v_after := jsonb_build_object(
      'serial', r.serial,
      'stock', v_stock_after,
      'l0_cost', r.l0_cost,
      'l0_conversion', 1,
      'l1_cost', r.l1_cost,
      'l1_relative_quantity', r.pcs_ctn,
      'l1_conversion', r.pcs_ctn,
      'fx_cfa', %(fx)s,
      'excel_tqty', r.tqty
    );

    INSERT INTO public.audit_log (
      actor_id, action, entity_type, entity_id, depot_id, before_value, after_value
    ) VALUES (
      NULL,
      'correct_electro2_stock_cost',
      'product',
      r.product_id,
      r.depot_id,
      v_before,
      v_after
    );

    v_done := v_done + 1;
  END LOOP;

  RAISE NOTICE 'ELECTRO 2 correction applied=%% skipped_idempotent=%%', v_done, v_skipped;
END $$;

-- After (verify)
WITH truth(serial, pcs_ctn, tqty, l0_cost, l1_cost) AS (
  VALUES
%(values)s
),
targets AS (
  SELECT
    t.serial,
    p.name AS product_name,
    p.current_stock_pieces AS stock,
    pu0.cost_price_cfa AS l0_cost,
    pu1.cost_price_cfa AS l1_cost,
    pu1.relative_quantity AS l1_rel,
    pu1.conversion_to_base AS l1_conv,
    EXISTS (
      SELECT 1 FROM public.audit_log a
      WHERE a.entity_id = p.id AND a.action = 'correct_electro2_stock_cost'
    ) AS corrected
  FROM truth t
  JOIN public.product_units pu
    ON lower(pu.serial_number) = lower(t.serial) AND pu.unit_level = 0
  JOIN public.products p ON p.id = pu.product_id
  LEFT JOIN public.product_units pu0 ON pu0.product_id = p.id AND pu0.unit_level = 0
  LEFT JOIN public.product_units pu1 ON pu1.product_id = p.id AND pu1.unit_level = 1
)
SELECT * FROM targets ORDER BY serial;
""" % {"n": len(rows), "fx": FX, "sv": sheet_value, "values": values}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(sql, encoding="utf-8")
    print("wrote", OUT, len(rows), "products", "value", round(sheet_value))

if __name__ == "__main__":
    main()
