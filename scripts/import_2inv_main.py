"""Import 2INV.xlsx catalog into Main Depot (Aquilux). One-shot admin RPC import."""
from __future__ import annotations

import json
import os
import re
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import urllib.error
import urllib.request
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "2INV.xlsx"
ENV = ROOT / ".env"

PLACEHOLDER_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
)

SERIAL_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._\-/]*$")
DIMENSION_RE = re.compile(r"\d+\s*[x×*]\s*\d+|cm\b|\d+m\b|velvet|^\d+[°?]?\s*$", re.I)


@dataclass
class ProductRow:
    sheet: str
    row: int
    serial: str
    name: str
    category: str  # divers | meubles
    cost: float | None
    sell: float | None
    tqty: float
    pcs_ctn: int = 1
    photos: list[tuple[bytes, str]] = field(default_factory=list)
    sell_pending: bool = False


def load_env() -> dict[str, str]:
    out: dict[str, str] = {}
    for line in ENV.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        out[k.strip()] = v.strip()
    return out


def as_float(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def clean_name(*parts: str | None) -> str:
    bits: list[str] = []
    for p in parts:
        if not p:
            continue
        t = " ".join(str(p).replace("\n", " ").split())
        if t:
            bits.append(t)
    return " — ".join(bits) if len(bits) > 1 else (bits[0] if bits else "")


def norm_serial(raw) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip().replace("`", "").replace("\n", " ")
    if not s or s.lower() in ("item no", "item no.", "ctn:"):
        return None
    if not SERIAL_RE.match(s):
        return None
    return s


def classify(sheet: str, serial: str, name: str) -> str:
    sn = serial.upper()
    text = name.lower()

    if sheet in ("RALONGE", "ELECRTRO", "Sheet8"):
        return "divers"

    if sheet in ("Sheet1", "Sheet2", "Sheet7"):
        return "meubles"

    if sheet == "Sheet3":
        if sn.startswith("491-"):
            return "meubles"
        if any(k in text for k in ("raincoat", "safety shoes", "safety helmet", "earplug", "cowhide")):
            return "meubles"
        return "divers"

    if sheet == "ELECTRO 2":
        if sn.startswith(("Z1-34", "Z1-35", "Z1-36", "Z1-80", "Z1-81")):
            return "meubles"
        if DIMENSION_RE.search(text) or DIMENSION_RE.search(serial):
            return "meubles"
        if sn.startswith("Z1-6") or sn.startswith("Z1-7") or sn.startswith("Z1-8") or sn.startswith("Z1-9"):
            if any(x in text for x in ("blender", "kettle", "2pin", "coffee", "air fryer", "hand blender")):
                return "divers"
            if not text or re.search(r"\d+\s*[x×*]", text):
                return "meubles"
        return "divers"

    return "meubles"


def image_bytes(img) -> tuple[bytes, str]:
    data = img._data()
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return data, "png"
    if data[:2] == b"\xff\xd8":
        return data, "jpg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return data, "webp"
    return PLACEHOLDER_PNG, "png"


def images_by_row(ws) -> dict[int, list[tuple[bytes, str]]]:
    out: dict[int, list[tuple[bytes, str]]] = defaultdict(list)
    for img in getattr(ws, "_images", []) or []:
        try:
            row = img.anchor._from.row + 1
        except Exception:
            continue
        out[row].append(image_bytes(img))
    return out


def parse_sheet1(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(7, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 1).value)
        if not serial:
            continue
        dim = ws.cell(r, 3).value
        color = ws.cell(r, 4).value or ws.cell(r, 5).value
        name = clean_name(str(dim) if dim else None, str(color) if color else None, serial)
        tqty = as_float(ws.cell(r, 8).value) or 0
        cost = as_float(ws.cell(r, 9).value)
        sell = as_float(ws.cell(r, 10).value)
        pcs = int(as_float(ws.cell(r, 7).value) or 1)
        rows.append(
            ProductRow(
                sheet="Sheet1",
                row=r,
                serial=serial,
                name=name or serial,
                category=classify("Sheet1", serial, name),
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=max(1, pcs),
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


def parse_sheet2(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(4, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 1).value)
        if not serial:
            continue
        name = clean_name(ws.cell(r, 3).value, ws.cell(r, 4).value, serial)
        tqty = as_float(ws.cell(r, 8).value) or 0
        cost = as_float(ws.cell(r, 15).value)
        sell = as_float(ws.cell(r, 16).value)
        pcs = int(as_float(ws.cell(r, 7).value) or 1)
        rows.append(
            ProductRow(
                sheet="Sheet2",
                row=r,
                serial=serial,
                name=name or serial,
                category=classify("Sheet2", serial, name),
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=max(1, pcs),
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


def parse_sheet3(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(3, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 2).value)
        if not serial:
            continue
        name = clean_name(ws.cell(r, 5).value, ws.cell(r, 4).value, serial)
        tqty = as_float(ws.cell(r, 8).value) or 0
        cost = as_float(ws.cell(r, 14).value)
        sell = as_float(ws.cell(r, 15).value)
        pcs = int(as_float(ws.cell(r, 7).value) or 1)
        rows.append(
            ProductRow(
                sheet="Sheet3",
                row=r,
                serial=serial,
                name=name or serial,
                category=classify("Sheet3", serial, name),
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=max(1, pcs),
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


def parse_ralonge(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(2, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 1).value)
        if not serial:
            continue
        name = clean_name(ws.cell(r, 3).value, ws.cell(r, 4).value, serial)
        tqty = as_float(ws.cell(r, 7).value) or 0
        cost = as_float(ws.cell(r, 13).value)
        sell = as_float(ws.cell(r, 14).value)
        rows.append(
            ProductRow(
                sheet="RALONGE",
                row=r,
                serial=serial,
                name=name or serial,
                category="divers",
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=1,
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


def parse_elecrtro(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(2, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 1).value)
        if not serial:
            continue
        name = clean_name(ws.cell(r, 3).value, ws.cell(r, 4).value, serial)
        tqty = as_float(ws.cell(r, 7).value) or 0
        cost = as_float(ws.cell(r, 13).value)
        sell = as_float(ws.cell(r, 14).value)
        rows.append(
            ProductRow(
                sheet="ELECRTRO",
                row=r,
                serial=serial,
                name=name or serial,
                category="divers",
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=1,
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


def parse_electro2(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(1, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 1).value)
        if not serial:
            continue
        name = clean_name(ws.cell(r, 3).value, serial)
        tqty = as_float(ws.cell(r, 8).value) or 0
        cost = as_float(ws.cell(r, 13).value)
        sell = as_float(ws.cell(r, 14).value)
        pcs = int(as_float(ws.cell(r, 7).value) or 1)
        cat = classify("ELECTRO 2", serial, name)
        rows.append(
            ProductRow(
                sheet="ELECTRO 2",
                row=r,
                serial=serial,
                name=name or serial,
                category=cat,
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=max(1, pcs),
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


def parse_sheet7(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(2, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 1).value)
        if not serial:
            continue
        name = clean_name(ws.cell(r, 3).value, serial)
        tqty = as_float(ws.cell(r, 6).value) or 0
        cost = as_float(ws.cell(r, 9).value)
        sell = as_float(ws.cell(r, 10).value)
        pcs = int(as_float(ws.cell(r, 5).value) or 1)
        rows.append(
            ProductRow(
                sheet="Sheet7",
                row=r,
                serial=serial,
                name=name or serial,
                category="meubles",
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=max(1, pcs),
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


def parse_sheet8(ws, imgs: dict[int, list]) -> list[ProductRow]:
    rows: list[ProductRow] = []
    for r in range(3, ws.max_row + 1):
        serial = norm_serial(ws.cell(r, 1).value)
        if not serial:
            continue
        tqty = as_float(ws.cell(r, 2).value) or 0
        cost = as_float(ws.cell(r, 3).value)
        sell = as_float(ws.cell(r, 4).value)
        label = {
            "STAB2000W": "Stabilizer 2000W",
            "STAB3000W": "Stabilizer 3000W",
            "KIIRIKO": "KIIRIKO",
            "FS30-1": "FS30-1",
        }.get(serial.upper().replace("`", ""), serial)
        rows.append(
            ProductRow(
                sheet="Sheet8",
                row=r,
                serial=serial,
                name=label,
                category="divers",
                cost=cost,
                sell=sell,
                tqty=tqty,
                pcs_ctn=1,
                photos=list(imgs.get(r, [])),
            )
        )
    return rows


PARSERS = {
    "Sheet1": parse_sheet1,
    "Sheet2": parse_sheet2,
    "Sheet3": parse_sheet3,
    "RALONGE": parse_ralonge,
    "ELECRTRO": parse_elecrtro,
    "ELECTRO 2": parse_electro2,
    "Sheet7": parse_sheet7,
    "Sheet8": parse_sheet8,
}


def parse_workbook(wb) -> list[ProductRow]:
    all_rows: list[ProductRow] = []
    for sheet_name, parser in PARSERS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        imgs = images_by_row(ws)
        all_rows.extend(parser(ws, imgs))

    used: set[str] = set()
    for rec in all_rows:
        key = rec.serial.lower()
        if key in used:
            rec.serial = f"{rec.serial}-{rec.sheet[:3].upper()}"
            key = rec.serial.lower()
        used.add(key)

        if rec.cost is not None and rec.sell is not None and rec.cost == rec.sell:
            rec.sell = 0.0
            rec.sell_pending = True
        elif rec.sell is None or rec.sell <= 0:
            rec.sell_pending = True
            if rec.sell is None:
                rec.sell = 0.0
        if rec.cost is None:
            rec.cost = 0.0

        rec.tqty = int(rec.tqty) if rec.tqty else 0

    return all_rows


MAX_PHOTO_BYTES = 5 * 1024 * 1024
COMPRESS_ABOVE_BYTES = 350 * 1024
MAX_EDGE = 1600


def compress_photo(content: bytes, ext: str) -> tuple[bytes, str]:
    if len(content) <= COMPRESS_ABOVE_BYTES and len(content) <= MAX_PHOTO_BYTES:
        return content, ext
    img = Image.open(BytesIO(content))
    img = img.convert("RGB")
    w, h = img.size
    scale = min(1.0, MAX_EDGE / float(max(w, h)))
    if scale < 1.0:
        img = img.resize(
            (max(1, round(w * scale)), max(1, round(h * scale))),
            Image.Resampling.LANCZOS,
        )
    quality = 82
    while quality >= 40:
        buf = BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        out = buf.getvalue()
        if len(out) <= MAX_PHOTO_BYTES:
            return out, "jpg"
        quality -= 10
    raise RuntimeError("photo still too large after compress")


class Api:
    def __init__(self, url: str, anon: str, token: str):
        self.url = url.rstrip("/")
        self.anon = anon
        self.token = token
        self._placeholder_url: str | None = None

    def _headers(self, extra=None):
        h = {
            "apikey": self.anon,
            "authorization": "Bearer %s" % self.token,
            "content-type": "application/json",
        }
        if extra:
            h.update(extra)
        return h

    def request(self, method, path, data=None, extra=None):
        body = None if data is None else json.dumps(data).encode("utf-8")
        req = urllib.request.Request(
            self.url + path,
            data=body,
            method=method,
            headers=self._headers(extra),
        )
        try:
            with urllib.request.urlopen(req) as resp:
                payload = resp.read()
                if not payload:
                    return None
                return json.loads(payload.decode("utf-8"))
        except urllib.error.HTTPError as e:
            err = e.read().decode("utf-8", errors="replace")
            raise RuntimeError("%s %s -> %s %s" % (method, path, e.code, err)) from e

    def rpc(self, name, params):
        return self.request("POST", "/rest/v1/rpc/%s" % name, params)

    def select(self, table, query):
        return self.request(
            "GET",
            "/rest/v1/%s?%s" % (table, query),
            extra={"accept": "application/json"},
        )

    def upload_photo(self, content: bytes, ext: str) -> str:
        content, ext = compress_photo(content, ext)
        name = "%s.%s" % (uuid.uuid4(), ext)
        content_type = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}[ext]
        req = urllib.request.Request(
            self.url + "/storage/v1/object/product-photos/" + name,
            data=content,
            method="POST",
            headers={
                "apikey": self.anon,
                "authorization": "Bearer %s" % self.token,
                "content-type": content_type,
                "x-upsert": "false",
            },
        )
        try:
            with urllib.request.urlopen(req) as resp:
                resp.read()
        except urllib.error.HTTPError as e:
            err = e.read().decode("utf-8", errors="replace")
            raise RuntimeError("upload %s -> %s %s" % (name, e.code, err)) from e
        return "%s/storage/v1/object/public/product-photos/%s" % (self.url, name)

    def photo_url_for(self, rec: ProductRow) -> str:
        if rec.photos:
            raw, ext = rec.photos[0]
            return self.upload_photo(raw, ext)
        if self._placeholder_url:
            return self._placeholder_url
        self._placeholder_url = self.upload_photo(PLACEHOLDER_PNG, "png")
        return self._placeholder_url


def login(url, anon, password):
    body = json.dumps({"email": "admin@vhm-pos.internal", "password": password}).encode("utf-8")
    req = urllib.request.Request(
        url.rstrip("/") + "/auth/v1/token?grant_type=password",
        data=body,
        method="POST",
        headers={"apikey": anon, "content-type": "application/json"},
    )
    with urllib.request.urlopen(req) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return data["access_token"]


def units_for(rec: ProductRow):
    cost = round(rec.cost or 0, 2)
    sell = round(rec.sell or 0, 2)
    pcs = int(rec.pcs_ctn or 1)
    piece = {
        "unit_level": 0,
        "unit_name": "Piece",
        "serial_number": rec.serial,
        "relative_quantity": None,
        "cost_price_cfa": cost,
        "selling_price_cfa": sell,
    }
    if pcs <= 1:
        return [piece]
    return [
        piece,
        {
            "unit_level": 1,
            "unit_name": "Carton",
            "serial_number": rec.serial + "-CTN",
            "relative_quantity": pcs,
            "cost_price_cfa": round(cost * pcs, 2),
            "selling_price_cfa": round(sell * pcs, 2),
        },
    ]


def product_id_from(created):
    if isinstance(created, dict) and "id" in created:
        return created["id"]
    if isinstance(created, list) and created and isinstance(created[0], dict):
        return created[0]["id"]
    raise RuntimeError("Unexpected create_product payload: %r" % created)


def ensure_category(api: Api, name: str) -> str:
    cats = api.select("categories", "select=id,name,is_active") or []
    hit = next((c for c in cats if str(c.get("name", "")).strip().lower() == name.lower()), None)
    if hit:
        return hit["id"]
    created = api.rpc("create_category", {"p_name": name.capitalize(), "p_parent_id": None})
    return created["id"]


def serial_to_product_info(api: Api) -> dict[str, dict]:
    """Map piece serial (lower) -> product_id + existing photo_url."""
    rows = (
        api.select(
            "product_units",
            "select=serial_number,product_id,products(photo_url)&unit_level=eq.0",
        )
        or []
    )
    out: dict[str, dict] = {}
    for row in rows:
        sn = str(row.get("serial_number", "")).strip().lower()
        if not sn:
            continue
        prod = row.get("products") or {}
        out[sn] = {
            "product_id": row["product_id"],
            "photo_url": prod.get("photo_url"),
        }
    return out


def main():
    env = load_env()
    url = env["VITE_SUPABASE_URL"]
    anon = env["VITE_SUPABASE_ANON_KEY"]

    wb = load_workbook(XLSX, data_only=True)
    products = parse_workbook(wb)
    no_photo = [p.serial for p in products if not p.photos]
    by_cat = defaultdict(int)
    for p in products:
        by_cat[p.category] += 1

    print("PARSED", len(products))
    print("BY_CATEGORY", dict(by_cat))
    print("NO_PHOTO", len(no_photo), "serials", no_photo[:20], "..." if len(no_photo) > 20 else "")
    print("SELL_PENDING", [p.serial for p in products if p.sell_pending])
    print("SUM_PIECES", sum(p.tqty for p in products))

    preview = ROOT / "scripts" / "_2inv_import_preview.json"
    preview.write_text(
        json.dumps(
            [
                {
                    "serial": p.serial,
                    "name": p.name,
                    "category": p.category,
                    "cost": p.cost,
                    "sell": p.sell,
                    "tqty": p.tqty,
                    "has_photo": bool(p.photos),
                    "sheet": p.sheet,
                }
                for p in products
            ],
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print("WROTE", preview)

    if os.environ.get("DRY_RUN") == "1":
        return

    password = os.environ.get("ADMIN_PASSWORD")
    if not password:
        raise SystemExit("Set ADMIN_PASSWORD (e.g. Admin2026!)")

    token = login(url, anon, password)
    api = Api(url, anon, token)

    depots = api.select("depots", "select=id,name,code,is_active&order=created_at")
    main_depot = next(
        (
            d
            for d in depots
            if str(d.get("code", "")).upper() in ("MAIN", "MAIN_DEPOT")
            or "main" in str(d.get("name", "")).lower()
        ),
        depots[0] if depots else None,
    )
    if not main_depot:
        raise SystemExit("No main depot")
    print("USING_DEPOT", main_depot)

    cat_ids = {
        "divers": ensure_category(api, "Divers"),
        "meubles": ensure_category(api, "Meubles"),
    }
    print("CATEGORIES", cat_ids)

    existing_units = api.select("product_units", "select=serial_number") or []
    existing_serials = {str(u["serial_number"]).lower() for u in existing_units}
    sync_existing = os.environ.get("SYNC_EXISTING") == "1"
    by_serial = serial_to_product_info(api) if sync_existing else {}

    report = []
    sync_report = []
    skipped = 0
    for rec in products:
        if rec.serial.lower() in existing_serials:
            if sync_existing and rec.serial.lower() in by_serial:
                info = by_serial[rec.serial.lower()]
                photo_url = info.get("photo_url") or api.photo_url_for(rec)
                api.rpc(
                    "update_product",
                    {
                        "p_product_id": info["product_id"],
                        "p_category_id": cat_ids[rec.category],
                        "p_name": rec.name,
                        "p_photo_url": photo_url,
                        "p_units": units_for(rec),
                        "p_track_low_stock": False,
                        "p_low_stock_threshold": None,
                    },
                )
                sync_report.append(
                    {
                        "serial": rec.serial,
                        "name": rec.name,
                        "category": rec.category,
                        "product_id": info["product_id"],
                        "cost": rec.cost,
                        "sell": rec.sell,
                        "has_photo": bool(rec.photos),
                    }
                )
                print("SYNC", rec.serial, rec.category)
            else:
                skipped += 1
                print("SKIP", rec.serial)
            continue
        photo_url = api.photo_url_for(rec)
        created = api.rpc(
            "create_product",
            {
                "p_depot_id": main_depot["id"],
                "p_category_id": cat_ids[rec.category],
                "p_name": rec.name,
                "p_photo_url": photo_url,
                "p_units": units_for(rec),
                "p_track_low_stock": False,
                "p_low_stock_threshold": None,
            },
        )
        pid = product_id_from(created)
        restock = None
        if rec.tqty > 0:
            restock = api.rpc(
                "restock_product",
                {
                    "p_product_id": pid,
                    "p_unit_level": 0,
                    "p_quantity_in_unit": rec.tqty,
                    "p_note": "Opening stock — 2INV.xlsx",
                },
            )
        existing_serials.add(rec.serial.lower())
        report.append(
            {
                "serial": rec.serial,
                "name": rec.name,
                "category": rec.category,
                "product_id": pid,
                "tqty": rec.tqty,
                "cost": rec.cost,
                "sell": rec.sell,
                "has_photo": bool(rec.photos),
                "stock": None if restock is None else restock.get("new_stock_pieces"),
                "sell_pending": rec.sell_pending,
            }
        )
        print("OK", rec.serial, rec.category, rec.tqty)

    out = ROOT / "scripts" / "_2inv_import_report.json"
    out.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    if sync_report:
        sync_out = ROOT / "scripts" / "_2inv_sync_report.json"
        sync_out.write_text(json.dumps(sync_report, indent=2, ensure_ascii=False), encoding="utf-8")
        print("SYNCED", len(sync_report), "wrote", sync_out)
    print("DONE new=%s skipped=%s wrote=%s" % (len(report), skipped, out))


if __name__ == "__main__":
    main()
